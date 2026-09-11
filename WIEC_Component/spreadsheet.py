"""Excel output for the WIEC_Component test items.

The column layout is deliberately identical to the one dune_hv_crate_test.py
produces, so both scripts keep appending to the same workbook:

    1-3    Test name / date / time
    4-5    Fan supply voltage / current
    6-11   Fan 1-6 RD
    12+    Fan PWM sweep   (num_fans * sweep_points * 4 columns)
    +8     TC1-4 resistance, TC1-4 temperature rise
    +96    HV, 12 columns per channel for 8 channels
"""

import os

import openpyxl


class ResultSpreadsheet:
    def __init__(self, json_data, prefix="WIEC Component"):
        self.json_data = json_data
        self.prefix = prefix
        self.rounding_factor = json_data["rounding_factor"]

        self.hv_cols = 12
        self.fan_rd_first_col = 6
        self.fan_pwm_first_col = self.fan_rd_first_col + 6
        self.fan_pwm_num_fans = int(json_data["keysight970a_fan_num"])
        self.fan_pwm_sweep_points = self.get_fan_pwm_sweep_points()
        self.fan_pwm_cols = (
            self.fan_pwm_num_fans * len(self.fan_pwm_sweep_points) * 4
        )
        self.tc_res_first_col = self.fan_pwm_first_col + self.fan_pwm_cols
        self.hv_res_first_col = self.tc_res_first_col + 8

        self.wb = None
        self.ws = None
        self.row = None
        self.path = None

    # ------------------------------------------------------------------
    # Fan PWM helpers -- shared by header creation and the actual sweep so
    # the header voltage labels always match the real sweep points.
    # ------------------------------------------------------------------
    def get_fan_pwm_sweep_points(self):
        steps = int(self.json_data["rigol832a_fan_voltage_number_of_data_samples"])
        fan_voltage_max = self.json_data["rigol832a_fan_voltage_max"]
        fan_voltage_min = self.json_data["rigol832a_fan_voltage_min"]

        # Keep spreadsheet construction from crashing before fan_pwm_sweep
        # gets a chance to print its config error message.
        if steps < 2:
            steps = 2

        prog_voltage = []
        for i in range(steps):
            voltage = (
                fan_voltage_min
                + (fan_voltage_max - fan_voltage_min) * i / (steps - 1)
            )
            prog_voltage.append(round(voltage, 3))

        prog_voltage[-1] = fan_voltage_max
        return prog_voltage

    @staticmethod
    def format_fan_pwm_voltage_label(voltage):
        try:
            return f"{float(voltage):g}"
        except (TypeError, ValueError):
            return str(voltage)

    def get_fan_pwm_headers(self, prog_voltage=None):
        if prog_voltage is None:
            prog_voltage = self.get_fan_pwm_sweep_points()

        headers = []
        for fan in range(1, self.fan_pwm_num_fans + 1):
            for supply_voltage in prog_voltage:
                label = self.format_fan_pwm_voltage_label(supply_voltage)
                headers.append(f"Fan #{fan} - {label} V : Supply Voltage [V]")
                headers.append(f"Fan #{fan} - {label} V : Read Voltage [V]")
                headers.append(f"Fan #{fan} - {label} V : Read Current [A]")
                headers.append(f"Fan #{fan} - {label} V : Oscillation Frequency [Hz]")
        return headers

    # ------------------------------------------------------------------
    # Workbook life cycle
    # ------------------------------------------------------------------
    def open(self, test_name, output_path):
        self.path = os.path.join(output_path, f"{self.json_data['output_file']}")

        if os.path.isfile(self.path):
            self.wb = openpyxl.load_workbook(filename=self.path)
            self.ws = self.wb.active
            self._ensure_styles()
            self.ensure_fan_pwm_columns()
            self.row = self.ws.max_row + 1
        else:
            self.wb = openpyxl.Workbook()
            self.ws = self.wb.active
            self.ws.title = self.json_data["sheet_name"]
            self._ensure_styles()
            self._build_headers()
            self.row = 3

        # The test name starts red, so a run that never finishes still reads
        # as a failure in the workbook.
        self.ws.cell(row=self.row, column=1, value=test_name).style = "fail"
        self.ws.cell(row=self.row, column=2, value=None)
        self.save()
        return self.row

    def _ensure_styles(self):
        """Add the named styles unless the workbook already carries them.

        Loading an old workbook that predates one of these styles used to make
        every '.style = "fail"' assignment raise, which lost the whole run.
        """
        existing = set(self.wb.named_styles)

        if "top" not in existing:
            top_style = openpyxl.styles.NamedStyle(name="top")
            top_style.font = openpyxl.styles.Font(bold=True, name="Calibri")
            bd = openpyxl.styles.Side(style="thin")
            top_style.border = openpyxl.styles.Border(
                left=bd, top=bd, right=bd, bottom=bd
            )
            self.wb.add_named_style(top_style)

        if "fail" not in existing:
            fail_style = openpyxl.styles.NamedStyle(name="fail")
            fail_style.font = openpyxl.styles.Font(color="FF0000")
            self.wb.add_named_style(fail_style)

        if "pass" not in existing:
            pass_style = openpyxl.styles.NamedStyle(name="pass")
            pass_style.font = openpyxl.styles.Font(color="009900")
            self.wb.add_named_style(pass_style)

    def _build_headers(self):
        ws = self.ws

        ws.cell(row=2, column=1, value="Test Name").style = "top"
        ws.cell(row=2, column=2, value="Date").style = "top"
        ws.cell(row=2, column=3, value="Time").style = "top"

        ws.cell(
            row=1,
            column=4,
            value="Fan Test - V/I for power supplying all 4 fans, RD signal outputs",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=4,
            end_row=1, end_column=self.fan_pwm_first_col - 1,
        )
        ws.cell(row=2, column=4, value="Supply Voltage").style = "top"
        ws.cell(row=2, column=5, value="Supply Current").style = "top"
        for i in range(1, 7):
            ws.cell(
                row=2, column=self.fan_rd_first_col - 1 + i, value=f"Fan {i} RD"
            ).style = "top"

        ws.cell(
            row=1,
            column=self.fan_pwm_first_col,
            value="Fan PWM Sweep - Programmed/read fan voltage, current, and oscillation frequency",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.fan_pwm_first_col,
            end_row=1, end_column=self.fan_pwm_first_col + self.fan_pwm_cols - 1,
        )
        for col_offset, header in enumerate(self.get_fan_pwm_headers()):
            ws.cell(
                row=2, column=self.fan_pwm_first_col + col_offset, value=header
            ).style = "top"

        for i in range(1, 5):
            ws.cell(
                row=2, column=self.tc_res_first_col - 1 + i, value=f"TC{i}_Resistance"
            ).style = "top"
            ws.cell(
                row=2, column=self.tc_res_first_col - 1 + 4 + i, value=f"TC{i}_Temp_Rise"
            ).style = "top"

        ws.cell(
            row=1,
            column=self.tc_res_first_col,
            value="Heater Test - Results for each heating element and temperature rise after heating time",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.tc_res_first_col,
            end_row=1, end_column=self.hv_res_first_col - 1,
        )

        ws.cell(
            row=1,
            column=self.hv_res_first_col,
            value="HV Test - Results for each configuration. Resistance in units shown, Setup Time k (seconds) in a*e^(-b * t)+c where k = 1/b ",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.hv_res_first_col,
            end_row=1, end_column=self.hv_res_first_col + 3 + (7 * self.hv_cols),
        )

        for i in range(8):
            base = self.hv_res_first_col + (i * self.hv_cols)
            for offset, label in enumerate([
                f"Ch{i}+ Open Res", f"Ch{i}+ Open Fit On", f"Ch{i}+ Open Fit Off",
                f"Ch{i}+ 10k Res", f"Ch{i}+ 10k Fit On", f"Ch{i}+ 10k Fit Off",
                f"Ch{i}- Open Res", f"Ch{i}- Open Fit On", f"Ch{i}- Open Fit Off",
                f"Ch{i}- 10k Res", f"Ch{i}- 10k Fit On", f"Ch{i}- 10k Fit Off",
            ]):
                ws.cell(row=2, column=base + offset, value=label).style = "top"

        column_letters = tuple(
            openpyxl.utils.get_column_letter(col + 1) for col in range(ws.max_column)
        )
        for column_letter in column_letters:
            ws.column_dimensions[column_letter].bestFit = True
        ws.freeze_panes = ws.cell(row=3, column=2)
        self.save()

    def ensure_fan_pwm_columns(self):
        """Insert (or rebuild) the one-row Fan PWM block in an older workbook."""
        ws = self.ws
        fan_pwm_headers = self.get_fan_pwm_headers()

        heater_col = None
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=1, column=col).value
            if isinstance(value, str) and value.startswith("Heater Test"):
                heater_col = col
                break

        expected_heater_col = self.fan_pwm_first_col + self.fan_pwm_cols
        already_has_new_fan_pwm = (
            ws.cell(row=2, column=self.fan_pwm_first_col).value == fan_pwm_headers[0]
        )
        heater_is_in_expected_place = heater_col == expected_heater_col

        # Drop stale merged ranges before moving columns around.
        for merged_range in list(ws.merged_cells.ranges):
            if merged_range.min_row == 1 and merged_range.max_row == 1:
                ws.unmerge_cells(str(merged_range))

        if not (already_has_new_fan_pwm and heater_is_in_expected_place):
            if heater_col is not None and heater_col > self.fan_pwm_first_col:
                ws.delete_cols(
                    self.fan_pwm_first_col, heater_col - self.fan_pwm_first_col
                )
            ws.insert_cols(self.fan_pwm_first_col, self.fan_pwm_cols)

        ws.cell(
            row=1,
            column=4,
            value="Fan Test - V/I for power supplying all 4 fans, RD signal outputs",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=4,
            end_row=1, end_column=self.fan_pwm_first_col - 1,
        )
        ws.cell(
            row=1,
            column=self.fan_pwm_first_col,
            value="Fan PWM Sweep - One-row programmed/read voltage, current, and oscillation frequency",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.fan_pwm_first_col,
            end_row=1, end_column=self.fan_pwm_first_col + self.fan_pwm_cols - 1,
        )
        ws.cell(
            row=1,
            column=self.tc_res_first_col,
            value="Heater Test - Results for each heating element and temperature rise after heating time",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.tc_res_first_col,
            end_row=1, end_column=self.hv_res_first_col - 1,
        )
        ws.cell(
            row=1,
            column=self.hv_res_first_col,
            value="HV Test - Results for each configuration. Resistance in units shown, Setup Time k (seconds) in a*e^(-b * t)+c where k = 1/b ",
        ).style = "top"
        ws.merge_cells(
            start_row=1, start_column=self.hv_res_first_col,
            end_row=1, end_column=self.hv_res_first_col + 3 + (7 * self.hv_cols),
        )

        for col_offset, header in enumerate(fan_pwm_headers):
            ws.cell(
                row=2, column=self.fan_pwm_first_col + col_offset, value=header
            ).style = "top"

    # ------------------------------------------------------------------
    # Writing
    # ------------------------------------------------------------------
    def write(self, column, value, passed=True, rounded=False):
        if rounded:
            try:
                value = round(float(value), self.rounding_factor)
            except (TypeError, ValueError):
                pass
        cell = self.ws.cell(row=self.row, column=column, value=value)
        if not passed:
            cell.style = "fail"
        return cell

    def write_stamp(self, date_string, time_string):
        self.ws.cell(row=self.row, column=2, value=date_string)
        self.ws.cell(row=self.row, column=3, value=time_string)
        self.save()

    def write_overall(self, test_name, passed):
        self.ws.cell(
            row=self.row, column=1, value=test_name
        ).style = "pass" if passed else "fail"
        self.save()

    def save(self):
        if self.wb is not None and self.path:
            self.wb.save(self.path)
